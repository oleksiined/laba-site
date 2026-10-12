#!/usr/bin/env python3
"""
LABA Catalog Importer
Читає Data_source_2.xlsx і завантажує дані в Supabase.

Використання:
  pip install openpyxl requests
  python scripts/import_catalog.py

Налаштування: задайте SUPABASE_URL і SUPABASE_KEY нижче
або через змінні оточення.
"""

import os
import sys
import json
import math
import openpyxl
import requests

# ── Конфігурація ──────────────────────────────────────────────
SUPABASE_URL = os.getenv("SUPABASE_URL", "https://XXXXXXXXXX.supabase.co")
SUPABASE_KEY = os.getenv("SUPABASE_KEY", "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...")
XLSX_PATH    = os.getenv("XLSX_PATH", "Data_source_2.xlsx")

HEADERS = {
    "apikey":        SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type":  "application/json",
    "Prefer":        "resolution=merge-duplicates",  # upsert
}

# ── Хелпери ───────────────────────────────────────────────────

def api(path, rows):
    """Upsert rows to Supabase table via REST."""
    if not rows:
        return
    url = f"{SUPABASE_URL}/rest/v1/{path}"
    # Upsert in batches of 200
    for i in range(0, len(rows), 200):
        batch = rows[i:i+200]
        r = requests.post(url, headers=HEADERS, json=batch)
        if r.status_code not in (200, 201):
            print(f"  ✗ {path} batch {i//200+1}: {r.status_code} {r.text[:200]}")
        else:
            print(f"  ✓ {path}: {len(batch)} рядків")

def safe(v):
    """Convert xlsx value to JSON-safe type."""
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s if s else None

def to_int(v):
    """Convert xlsx numeric to int, or None."""
    v = safe(v)
    if v is None:
        return None
    try:
        return int(float(str(v)))
    except (ValueError, TypeError):
        return None

def to_price(v):
    """Convert price string/number to numeric or None."""
    v = safe(v)
    if v is None:
        return None
    try:
        f = float(str(v).replace(',', '.').replace('\xa0', '').strip())
        return None if math.isnan(f) else f
    except (ValueError, TypeError):
        return None

# Відповідність turnaround_unit → дні
UNIT_TO_DAYS = {
    "роб. днів": 1, "робочих днів": 1,
    "calendar days": 1,
    "год": 0,  # hours, skip
}

def turnaround_days(value, unit):
    value = to_int(value)
    if not value:
        return None
    unit = str(unit or '').strip().lower()
    if 'год' in unit or 'hour' in unit:
        return None  # hours — not useful for display
    return value  # working days


# ── Завантаження xlsx ─────────────────────────────────────────

def load_sheet(wb, name, skip_rows=2):
    """Return list of dicts for a sheet, skipping header + description rows."""
    ws = wb[name]
    headers = [str(c.value).strip() if c.value else f"col{i}" for i, c in enumerate(ws[1])]
    rows = []
    for row in ws.iter_rows(min_row=skip_rows + 1, values_only=True):
        if not any(row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def run():
    if not os.path.exists(XLSX_PATH):
        print(f"✗ Файл не знайдено: {XLSX_PATH}")
        sys.exit(1)

    print(f"📖 Читаємо {XLSX_PATH}…")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    # ── 1. Materials ─────────────────────────────────────────
    print("\n→ Materials")
    mat_rows = load_sheet(wb, "Materials")
    materials = []
    for r in mat_rows:
        mid = safe(r.get("material_id"))
        if not mid:
            continue
        materials.append({
            "material_id":   mid,
            "material_code": safe(r.get("material_code")),
            "name":          safe(r.get("name")) or mid,
            "short_name":    safe(r.get("short_name")),
            "sort_order":    to_int(r.get("sort_order")),
            "status":        safe(r.get("status")) or "active",
        })
    api("materials", materials)

    # ── 2. Services ──────────────────────────────────────────
    print("\n→ Services")
    svc_rows = load_sheet(wb, "Services")
    services = []
    for r in svc_rows:
        sid = safe(r.get("service_id"))
        if not sid:
            continue
        status = safe(r.get("status")) or "draft"
        if status not in ("active", "draft"):
            continue  # skip inactive
        td_val  = r.get("turnaround_value")
        td_unit = r.get("turnaround_unit")
        services.append({
            "service_id":          sid,
            "material_id":         safe(r.get("material_id")),
            "service_code":        safe(r.get("service_code")),
            "group_name":          safe(r.get("group_name")),
            "name":                safe(r.get("name")) or sid,
            "short_name":          safe(r.get("short_name")),
            "symbol":              safe(r.get("symbol")),
            "short_description":   safe(r.get("short_description")),
            "price":               to_price(r.get("price")),
            "currency":            safe(r.get("currency")) or "UAH",
            "turnaround_days":     turnaround_days(td_val, td_unit),
            "status":              status,
            "sort_order":          to_int(r.get("sort_order")),
            "orderable_standalone": bool(r.get("orderable_standalone")),
        })
    api("services", services)

    # ── 3. Packages ──────────────────────────────────────────
    print("\n→ Packages")
    pkg_rows = load_sheet(wb, "Packages")
    packages = []
    for r in pkg_rows:
        pid = safe(r.get("package_id"))
        if not pid:
            continue
        status = safe(r.get("status")) or "draft"
        td_val  = r.get("turnaround_value")
        td_unit = r.get("turnaround_unit")
        packages.append({
            "package_id":        pid,
            "material_id":       safe(r.get("material_id")),
            "package_group":     safe(r.get("package_group")),
            "name":              safe(r.get("name")) or pid,
            "short_description": safe(r.get("short_description")),
            "price":             to_price(r.get("price")),
            "currency":          safe(r.get("currency")) or "UAH",
            "turnaround_days":   turnaround_days(td_val, td_unit),
            "status":            status,
            "sort_order":        to_int(r.get("sort_order")),
            "slug":              safe(r.get("slug")),
        })
    api("packages", packages)

    # ── 4. Package Items ─────────────────────────────────────
    print("\n→ Package Items")
    pi_rows = load_sheet(wb, "Package Items")
    items = []
    for r in pi_rows:
        pid = safe(r.get("package_id"))
        sid = safe(r.get("service_id"))
        if not pid or not sid:
            continue
        items.append({
            "package_id":    pid,
            "service_id":    sid,
            "display_order": to_int(r.get("display_order")) or 0,
            "required":      bool(r.get("required") is not False),
        })
    api("package_items", items)

    print("\n✅ Імпорт завершено!")
    print(f"   Матеріалів: {len(materials)}")
    print(f"   Послуг:     {len(services)}")
    print(f"   Пакетів:    {len(packages)}")
    print(f"   Позицій:    {len(items)}")


if __name__ == "__main__":
    run()
